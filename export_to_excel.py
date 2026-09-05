import json
import sys
from openpyxl import Workbook


def export_to_excel(input_file, output_file):
    # Read query data
    with open(input_file, "r", encoding="utf-8") as file:
        data = json.load(file)

    columns = data["columns"]
    rows = data["rows"]

    # Create workbook
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Query Results"

    # Write column headers
    for column_number, column_name in enumerate(columns, start=1):
        worksheet.cell(
            row=1,
            column=column_number,
            value=column_name
        )

    # Write query results
    for row_number, row_data in enumerate(rows, start=2):
        for column_number, value in enumerate(row_data, start=1):
            worksheet.cell(
                row=row_number,
                column=column_number,
                value=value
            )

    # Adjust column widths
    for column in worksheet.columns:
        max_length = 0

        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        worksheet.column_dimensions[
            column[0].column_letter
        ].width = min(max_length + 2, 50)

    workbook.save(output_file)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python3 export_to_excel.py <input.json> <output.xlsx>")
        sys.exit(1)

    export_to_excel(sys.argv[1], sys.argv[2])

    print(f"Excel file created: {sys.argv[2]}")
